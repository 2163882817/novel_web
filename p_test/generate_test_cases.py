from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).with_name("测试用例.xlsx")
HEADERS = ["用例编号", "模块", "优先级", "测试类型", "前置条件", "测试步骤", "预期结果", "自动化状态", "备注"]
CASES = [
    ["F-01", "服务可用性", "P0", "功能", "前后端已启动", "打开书架并调用健康检查", "书架可加载；健康检查返回正常状态", "待自动化", "冒烟"],
    ["F-02", "新建小说", "P0", "功能", "位于空书架", "填写书名后创建", "创建成功并进入对应工作台", "已自动化", "mock API"],
    ["F-03", "新建小说", "P0", "功能", "已打开创建弹窗", "不填写书名直接创建", "阻止提交并提示请填写书名", "已自动化", "前端校验"],
    ["F-04", "小说隔离", "P0", "功能", "存在两部小说", "分别创建章节、人物、伏笔", "数据互不混入", "待自动化", "数据隔离"],
    ["F-05", "章节保存", "P0", "功能", "已进入工作台", "填写正文并保存章节", "标题、正文、字数、序号正确保存", "待自动化", "持久化"],
    ["F-06", "章节编辑", "P0", "功能", "存在章节", "修改标题、正文、细纲并刷新", "刷新后显示修改后的内容", "待自动化", "持久化"],
    ["F-07", "章节删除", "P0", "功能", "存在章节", "分别取消/确认删除", "取消不变；确认后章节消失", "待自动化", "高风险操作"],
    ["F-08", "流式写作", "P0", "功能", "已填写细纲，SSE 已模拟", "开始生成并等待完成", "增量写入；完成后恢复操作状态", "待自动化", "SSE"],
    ["F-09", "流式写作", "P0", "功能", "生成进行中", "点击停止", "停止后不再追加内容", "待自动化", "AbortController"],
    ["F-10", "导出 TXT", "P0", "功能", "存在章节", "勾选章节并导出", "下载成功，文件名与正文内容正确", "已自动化", "mock 下载"],
    ["F-11", "API 配置", "P0", "安全/功能", "已有保存的密钥", "打开配置页并保存空 Key", "仅显示尾号；不回显完整 Key；保留原 Key", "已自动化", "脱敏"],
    ["F-12", "错误处理", "P0", "功能", "后端返回 404", "访问不存在小说工作台", "显示清晰错误，不出现无限加载", "已自动化", "mock 404"],
    ["F-13", "书架", "P1", "UI", "分别准备空/加载中/错误/多卡片数据", "打开书架", "四种状态文案和布局正确", "待自动化", "视觉回归"],
    ["F-14", "AI 细纲", "P1", "功能", "模拟细纲服务", "点击生成细纲", "细纲和推荐标题写入编辑区", "待自动化", "mock AI"],
    ["F-15", "A/B/C 变体", "P1", "功能", "已填写细纲", "生成变体并导入主编辑器", "三版展示正确，导入后正文更新", "待自动化", "mock AI"],
    ["F-16", "标题生成", "P1", "功能", "有章节正文", "生成候选标题并选择", "标题输入框更新", "待自动化", "mock AI"],
    ["F-17", "校对修订", "P1", "功能", "存在已保存章节", "校对、取消选择、修订、应用预览", "选择规则正确，正文在应用后更新", "待自动化", "mock AI"],
    ["F-18", "定稿记忆", "P1", "功能", "存在章节与记忆数据", "执行定稿", "状态、摘要和记忆按结果刷新", "待自动化", "mock AI"],
    ["F-19", "人物卡", "P1", "功能", "进入设定页", "新建、编辑、删除人物卡", "字段和列表正确更新", "待自动化", "CRUD"],
    ["F-20", "伏笔库", "P1", "功能", "进入设定页", "新建、编辑、删除伏笔并切换状态", "描述与状态正确保存", "待自动化", "CRUD"],
    ["F-21", "卷管理", "P1", "功能", "进入设定页", "新建卷、编辑大纲、生成摘要", "编号、状态和内容正确", "待自动化", "mock AI"],
    ["F-22", "文档导入", "P1", "功能", "进入导入 Tab", "导入故事圣经/总纲/人物卡/限制词", "replace/append 语义正确，非法格式有提示", "待自动化", "文件上传"],
    ["U-01", "全局布局", "P1", "UI", "固定浏览器和测试数据", "打开四个主要页面", "标题、导航、主操作和内容区层级清晰", "待自动化", "视觉基线"],
    ["U-02", "书架卡片", "P1", "UI", "存在长标题小说", "检查卡片与删除按钮", "无覆盖、溢出或误触", "待自动化", "1440px"],
    ["U-03", "创建弹窗", "P1", "UI/可访问性", "打开创建弹窗", "检查遮罩、焦点、取消与提交", "弹窗可操作，背景不误触", "待自动化", "键盘"],
    ["U-04", "工作台", "P1", "UI", "长正文和多章节数据", "检查章节列表和编辑器", "可滚动、无溢出、操作区清晰", "待自动化", "视觉基线"],
    ["U-05", "AI 状态", "P1", "UI", "模拟成功/失败/生成中", "触发 AI 操作", "加载、禁用、成功、错误反馈清晰", "待自动化", "异步状态"],
    ["U-06", "导出弹窗", "P1", "UI", "存在多个章节", "全选、清空、勾选、导出", "按钮状态与勾选状态一致", "待自动化", "视觉基线"],
    ["U-07", "设定页 Tab", "P1", "UI", "进入设定页", "切换人物、伏笔、卷、导入", "激活态清晰，不残留错误表单状态", "待自动化", "视觉基线"],
    ["U-08", "配置页", "P1", "UI/安全", "已有保存配置", "检查掩码、提示与对比度", "不泄露 Key，反馈文案清晰", "待自动化", "安全"],
    ["U-09", "响应式布局", "P1", "UI", "设置 1280×800、768×1024、390×844", "访问核心页面并操作弹窗", "无横向溢出，按钮与输入框可用", "待自动化", "多视口"],
    ["U-10", "中文显示", "P1", "UI", "模拟正常与错误响应", "检查页面、提示、下载文件名", "无乱码、截断和占位符残留", "待自动化", "本地化"],
]

def build_sheet(workbook):
    sheet = workbook.active
    sheet.title = "测试用例"
    sheet.append(HEADERS)
    for row in CASES:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    widths = [12, 14, 10, 14, 22, 34, 34, 14, 18]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + column)].width = width
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 48
    priority_validation = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=False)
    status_validation = DataValidation(type="list", formula1='"待自动化,已自动化,不自动化,阻塞"', allow_blank=False)
    sheet.add_data_validation(priority_validation)
    sheet.add_data_validation(status_validation)
    priority_validation.add(f"C2:C{sheet.max_row}")
    status_validation.add(f"H2:H{sheet.max_row}")
    sheet.conditional_formatting.add(f"C2:C{sheet.max_row}", FormulaRule(formula=["C2=\"P0\""], fill=PatternFill("solid", fgColor="F4CCCC")))
    sheet.conditional_formatting.add(f"H2:H{sheet.max_row}", FormulaRule(formula=["H2=\"已自动化\""], fill=PatternFill("solid", fgColor="D9EAD3")))

def build_summary(workbook):
    sheet = workbook.create_sheet("执行概览")
    sheet.append(["指标", "数量"])
    sheet.append(["用例总数", len(CASES)])
    sheet.append(["P0 用例", sum(case[2] == "P0" for case in CASES)])
    sheet.append(["P1 用例", sum(case[2] == "P1" for case in CASES)])
    sheet.append(["已自动化", sum(case[7] == "已自动化" for case in CASES)])
    sheet.append(["待自动化", sum(case[7] == "待自动化" for case in CASES)])
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 14

workbook = Workbook()
build_sheet(workbook)
build_summary(workbook)
workbook.save(OUTPUT)
print(f"已生成：{OUTPUT}")
